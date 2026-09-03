# -*- coding: utf-8 -*-
"""네이버 키워드도구 다운로드 파일 합산 -> 게이트 1 판정.

입력: data/keywords/gate1/*.xlsx  (키워드도구 '다운로드' 산출물)
출력: data/keywords/gate1-volumes.csv + 판정 요약

게이트 1 기준 (CLAUDE.md): 상위 100개 키워드 월간 검색량 합계 10만 미만이면 중단.
'<10'은 5로 근사한다. 여러 파일에 같은 키워드가 나오면 최대값만 센다.
"""
import csv
import glob
import os
import re

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "keywords", "gate1", "*.xlsx")
ITEMS = os.path.join(ROOT, "data", "keywords", "items.csv")
OUT = os.path.join(ROOT, "data", "keywords", "gate1-volumes.csv")
GATE = 100000

# 우리 사이트가 답할 수 있는 검색어인지 가르는 축
DISPOSAL = re.compile(r"(버리|분리수거|분리배출|폐기|배출|음식물|종량제|재활용)")


def volume(cell):
    if cell is None:
        return 0
    s = str(cell).strip().replace(",", "")
    if s.startswith("<"):
        return 5
    return int(s) if s.isdigit() else 0


def load_items():
    with open(ITEMS, encoding="utf-8-sig") as f:
        out = set()
        for r in csv.DictReader(f):
            out.add(re.sub(r"\s+", "", r["item"]))
            for v in filter(None, r["variants"].split("|")):
                out.add(re.sub(r"\s+", "", v))
    return out


def main():
    items = load_items()
    best = {}
    files = sorted(glob.glob(SRC))
    for path in files:
        ws = openpyxl.load_workbook(path, read_only=True)[
            openpyxl.load_workbook(path, read_only=True).sheetnames[0]]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 2 or not row[0]:
                continue
            kw = str(row[0]).strip()
            best[kw] = max(best.get(kw, 0), volume(row[1]) + volume(row[2]))

    rows = []
    for kw, vol in best.items():
        matched = next((it for it in items if it in kw), "")
        rows.append({
            "keyword": kw, "volume": vol,
            "item": matched,
            "on_axis": 1 if (matched and DISPOSAL.search(kw)) else 0,
        })
    rows.sort(key=lambda r: -r["volume"])

    with open(OUT, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["keyword", "volume", "item", "on_axis"])
        w.writeheader()
        w.writerows(rows)

    axis = [r for r in rows if r["on_axis"]]
    loose = sum(r["volume"] for r in rows[:100])
    strict = sum(r["volume"] for r in axis[:100])

    print("파일 %d개 / 키워드 %d개 (품목축 %d개)" % (len(files), len(rows), len(axis)))
    print()
    print("느슨한 집계  상위 100개 합계: %10s" % format(loose, ","))
    print("엄격한 집계  상위 100개 합계: %10s   <- 판정 기준" % format(strict, ","))
    print("기준선: %s" % format(GATE, ","))
    print()
    print("판정: %s" % ("통과 (진행)" if strict >= GATE else "미달 (추가 수집 필요)"))
    print()
    print("품목축 상위 15:")
    for r in axis[:15]:
        print("  %-26s %9s  [%s]" % (r["keyword"][:26], format(r["volume"], ","), r["item"]))
    print("\n-> %s" % OUT)


if __name__ == "__main__":
    main()
